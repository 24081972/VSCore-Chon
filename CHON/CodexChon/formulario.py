import tkinter as tk
from tkinter import messagebox
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


ARCHIVO_EXCEL = Path(__file__).with_name("datos_formulario.xlsx")


def guardar_en_excel(nombre, correo, edad, comentario):
    if ARCHIVO_EXCEL.exists():
        libro = load_workbook(ARCHIVO_EXCEL)
        hoja = libro.active
    else:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Datos"
        hoja.append(["Fecha", "Nombre", "Correo", "Edad", "Comentario"])

    hoja.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        nombre,
        correo,
        int(edad),
        comentario if comentario else "Sin comentario",
    ])
    libro.save(ARCHIVO_EXCEL)


def enviar_formulario():
    nombre = entrada_nombre.get().strip()
    correo = entrada_correo.get().strip()
    edad = entrada_edad.get().strip()
    comentario = texto_comentario.get("1.0", tk.END).strip()

    if not nombre or not correo or not edad:
        messagebox.showwarning("Datos incompletos", "Completa nombre, correo y edad.")
        return

    if not edad.isdigit():
        messagebox.showwarning("Edad invalida", "La edad debe ser un numero.")
        return

    guardar_en_excel(nombre, correo, edad, comentario)

    messagebox.showinfo(
        "Enviado",
        f"Los datos se guardaron correctamente en:\n{ARCHIVO_EXCEL}",
    )
    limpiar_formulario()


def limpiar_formulario():
    entrada_nombre.delete(0, tk.END)
    entrada_correo.delete(0, tk.END)
    entrada_edad.delete(0, tk.END)
    texto_comentario.delete("1.0", tk.END)
    entrada_nombre.focus()


ventana = tk.Tk()
ventana.title("Formulario")
ventana.geometry("420x360")
ventana.resizable(False, False)

contenedor = tk.Frame(ventana, padx=20, pady=20)
contenedor.pack(fill="both", expand=True)

tk.Label(contenedor, text="Formulario de registro", font=("Arial", 16, "bold")).grid(
    row=0, column=0, columnspan=2, pady=(0, 18)
)

tk.Label(contenedor, text="Nombre:").grid(row=1, column=0, sticky="w", pady=5)
entrada_nombre = tk.Entry(contenedor, width=32)
entrada_nombre.grid(row=1, column=1, pady=5)

tk.Label(contenedor, text="Correo:").grid(row=2, column=0, sticky="w", pady=5)
entrada_correo = tk.Entry(contenedor, width=32)
entrada_correo.grid(row=2, column=1, pady=5)

tk.Label(contenedor, text="Edad:").grid(row=3, column=0, sticky="w", pady=5)
entrada_edad = tk.Entry(contenedor, width=32)
entrada_edad.grid(row=3, column=1, pady=5)

tk.Label(contenedor, text="Comentario:").grid(row=4, column=0, sticky="nw", pady=5)
texto_comentario = tk.Text(contenedor, width=24, height=5)
texto_comentario.grid(row=4, column=1, pady=5)

botones = tk.Frame(contenedor)
botones.grid(row=5, column=0, columnspan=2, pady=18)

tk.Button(botones, text="Enviar", width=12, command=enviar_formulario).pack(
    side="left", padx=6
)
tk.Button(botones, text="Limpiar", width=12, command=limpiar_formulario).pack(
    side="left", padx=6
)
tk.Button(botones, text="Salir", width=12, command=ventana.destroy).pack(
    side="left", padx=6
)

entrada_nombre.focus()
ventana.mainloop()
